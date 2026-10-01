# -*- coding: utf-8 -*-
import csv
import json
import re
import openpyxl

CATEGORY_STRUCTURE = {
    "Embutidos": {
        "name": "Embutidos",
        "icon": "Beef",
        "color": "#e11d48",
        "subcategories": ["Jamones", "Ahumados", "Chorizos", "Mortadelas", "Salchichas"]
    },
    "Bebidas": {
        "name": "Bebidas",
        "icon": "Coffee",
        "color": "#f59e0b",
        "subcategories": ["Jugos y Néctares", "Light", "Té"]
    },
    "Leche Líquida": {
        "name": "Leche Líquida",
        "icon": "Milk",
        "color": "#0284c7",
        "subcategories": ["Líquida UHT"]
    },
    "Leche en Polvo": {
        "name": "Leche en Polvo",
        "icon": "Package",
        "color": "#8b5cf6",
        "subcategories": ["Polvo"]
    },
    "Derivados Lácteos": {
        "name": "Derivados Lácteos",
        "icon": "Utensils",
        "color": "#10b981",
        "subcategories": ["Crema de Leche", "Mantequilla", "Quesos de Cabra"]
    },
    "Huevos": {
        "name": "Huevos",
        "icon": "Egg",
        "color": "#f97316",
        "subcategories": ["Huevos"]
    }
}

def classify(name, seg=''):
    n = (name or '').upper().strip()
    s = (seg or '').upper().strip()

    # 1. HUEVOS
    if 'HUEVO' in n or 'PUROVO' in n or s == 'PRODHUEVOS':
        return ('Huevos', 'Huevos')

    # 2. EMBUTIDOS
    if 'SALCHICHA' in n:
        return ('Embutidos', 'Salchichas')
    if 'MORTADELA' in n:
        return ('Embutidos', 'Mortadelas')
    if 'CHORIZO' in n or 'CHORIZON' in n or 'CHORIZÓN' in n:
        return ('Embutidos', 'Chorizos')
    if any(k in n for k in ['AHUMADO', 'AHUMADA', 'TOCINETA', 'CHULETA']):
        return ('Embutidos', 'Ahumados')
    if any(k in n for k in ['JAMON', 'JAMÓN', 'FIAMBRE', 'PIERNA', 'ESPALDA', 'PECHUGA', 'VISKING', 'AREPERO', 'ALIMEX', 'ITALICO', 'PLUMROSE', 'HERMO']) or s in ['PRODJAMONERIA', 'PRODEMBUTIDOS']:
        return ('Embutidos', 'Jamones')

    # 3. DERIVADOS LACTEOS
    if ('CREMA DE LECHE' in n or re.search(r'\bCREMA\b', n) or 'SUERO' in n or 'NATA' in n) and 'DESCREMADA' not in n:
        return ('Derivados Lácteos', 'Crema de Leche')
    if 'MANTEQUILLA' in n:
        return ('Derivados Lácteos', 'Mantequilla')
    if any(k in n for k in ['CABRA', 'QUESO', 'ANANKE', 'ANANKÉ', 'QUIBORENO', 'QUIBOREÑO', 'QUBORENO', 'MADURADO', 'SEMIMADURADO', 'GOUDA', 'PARMESANO', 'EDAM', 'TILLITA', 'PAISA BLANCO']):
        return ('Derivados Lácteos', 'Quesos de Cabra')

    # 4. LECHE EN POLVO
    if any(k in n for k in ['LECHE EN POLVO', 'POLV.', 'POLVO', 'PURISIMA 900GR', 'PURISIMA 400GR', 'PURISIMA 200GR', 'PURISIMA 500GR', 'PURISIMA 1KG']) and not any(k in n for k in ['LIQUIDA', 'BARISTA', 'UHT']):
        return ('Leche en Polvo', 'Polvo')
    if 'CAMPINA' in n or 'CAMPIÑA' in n:
        if any(g in n for g in ['400', '900', '200', '1KG', 'SOBRE', 'POLVO']):
            return ('Leche en Polvo', 'Polvo')

    # 5. LECHE LIQUIDA
    if any(k in n for k in ['LECHE COMPLETA', 'LECHE DESCREMADA', 'LECHE  DESLACTOSADA', 'LECHE DESLACTOSADA', 'LECHE SEMIDESCREMADA', 'BARISTA LIQUIDA', 'BARISTA', 'LECHE UHT', 'UHT']) or s == 'PRODBEBIDAS_LACT' or (('LECHE' in n) and any(m in n for m in ['PASTORENA', 'NATULAC', 'SAN SIMON', 'PURISIMA']) and 'POLVO' not in n and 'POLV' not in n and '900' not in n and '400' not in n):
        return ('Leche Líquida', 'Líquida UHT')

    # 6. BEBIDAS
    if re.search(r'\bT[EÉ]\b', n) or ('TE ' in n or n.startswith('TE ') or 'TÉ ' in n or n.startswith('TÉ ')):
        return ('Bebidas', 'Té')
    if 'LIGHT' in n or 'LIGTH' in n:
        return ('Bebidas', 'Light')
    if any(k in n for k in ['NECTAR', 'NÉCTAR', 'JUGO', 'NARANJADA', 'DURAZNO', 'MANZANA', 'PERA', 'NARANJA', 'FRUTAS', 'FRICA', 'KAITO', 'YUKERY']) or s == 'PRODBEBIDAS':
        return ('Bebidas', 'Jugos y Néctares')

    return ('Embutidos', 'Jamones')

def extract_competitor_brand(name):
    n = (name or '').upper()
    if 'NATULAC' in n: return 'Natulac'
    if 'PASTORENA' in n or 'PASTOREÑA' in n: return 'Pastoreña'
    if 'SAN SIMON' in n or 'SAN SIMÓN' in n: return 'San Simón'
    if 'CAMPINA' in n or 'CAMPIÑA' in n: return 'La Campiña'
    if 'PLUMROSE' in n: return 'Plumrose'
    if 'PAISA' in n: return 'Paisa'
    if 'FRICA' in n: return 'Frica'
    if 'KAITO' in n: return 'Kaito'
    if 'PARMALAT' in n: return 'Parmalat'
    if 'FIESTA' in n: return 'Fiesta'
    if 'HERMO' in n: return 'Hermo'
    if 'TORONDOI' in n or 'TORONDOY' in n: return 'Torondoy'
    if 'MARACAY' in n or 'LACTEOS MARACAY' in n: return 'Lácteos Maracay'
    if 'MONSERRATINA' in n or 'MONTSERRATINA' in n: return 'La Montserratina'
    if 'MILLENIUM' in n: return 'Millenium'
    if 'ANANKE' in n or 'ANANKÉ' in n: return 'Ananké'
    if 'QUIBORENO' in n or 'QUIBOREÑO' in n or 'QUBORENO' in n: return 'Quiboreño'
    if 'PUROVO' in n: return 'Purovo'
    if 'RICCI' in n: return 'Ricci'
    if 'LEYTON' in n: return 'Leyton'
    if 'QUENACA' in n: return 'Quenaca'
    if 'YUKERY' in n: return 'Yukery'
    if 'OSCAR MAYER' in n: return 'Oscar Mayer'
    if 'FLOR DE ARAGUA' in n: return 'Flor de Aragua'
    return 'Otra Competencia'

def extract_tunal_brand(name):
    n = (name or '').upper()
    if 'PURISIMA' in n: return 'Purísima'
    if 'ALIMEX' in n: return 'Alimex'
    if 'ITALICO' in n or 'ITÁLICO' in n: return 'Itálico'
    return 'El Tunal'

def extract_presentation(name):
    n = (name or '').upper()
    if '1L' in n or '1 L' in n or '1 LITRO' in n or '1LT' in n: return '1 Litro'
    if '900GR' in n or '900 GR' in n or '900G' in n: return '900g'
    if '400GR' in n or '400 GR' in n or '400G' in n: return '400g'
    if '500GR' in n or '500 GR' in n or '500G' in n or '500 ML' in n or '500ML' in n or '50OML' in n: return '500g/ml'
    if '250GR' in n or '250 GR' in n or '250G' in n or '250 ML' in n or '250ML' in n: return '250g/ml'
    if '200GR' in n or '200 GR' in n or '200G' in n or '200 ML' in n or '200ML' in n: return '200g/ml'
    if '225GR' in n or '225 GR' in n or '225G' in n: return '225g'
    if '450GR' in n or '450 GR' in n or '450G' in n: return '450g'
    if '800GR' in n or '800 GR' in n or '800G' in n: return '800g'
    if '1KG' in n or '1 KG' in n or '1000G' in n: return '1kg'
    if '30 UND' in n or '30UND' in n or '30 UNIDADES' in n: return '30 und'
    if '15 UND' in n or '15UND' in n or '1X15' in n: return '15 und'
    if '12' in n or '1X12' in n: return '12 und'
    if '6' in n or '1X6' in n or '2 X 6' in n: return '6 und'
    return 'Estándar'

def load_catalog_from_excel(filepath):
    catalog = []
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        for sheet_name in wb.sheetnames:
            if sheet_name == 'Portafolio':
                continue
            ws = wb[sheet_name]
            current_subcat = ""
            for r in range(4, ws.max_row + 1):
                cat = ws.cell(r, 2).value
                sub = ws.cell(r, 3).value
                if sub:
                    current_subcat = str(sub).strip()
                marca = ws.cell(r, 4).value
                code = ws.cell(r, 5).value
                desc = ws.cell(r, 6).value
                competidores = ws.cell(r, 8).value
                
                if desc:
                    cat_name = str(cat).strip() if cat else sheet_name
                    if 'Leche UHT' in cat_name or 'Leche UHT' in sheet_name:
                        cat_name = 'Leche Líquida'
                    elif 'Leche Polvo' in cat_name or 'Leche Polvo' in sheet_name:
                        cat_name = 'Leche en Polvo'
                    elif 'Derivados' in cat_name or 'Derivados' in sheet_name:
                        cat_name = 'Derivados Lácteos'

                    catalog.append({
                        'sheet': sheet_name,
                        'categoria': cat_name,
                        'subcategoria': current_subcat,
                        'marca': str(marca).strip() if marca else '',
                        'codigo': str(code).strip() if code else '',
                        'descripcion': str(desc).strip(),
                        'competidores': str(competidores).strip() if competidores else ''
                    })
    except Exception as e:
        print("Error reading Excel:", e)
    return catalog

def main():
    print("Processing datasets...")
    catalog = load_catalog_from_excel('data/Portafolio y Competencia (1).xlsx')
    print(f"Catalog items extracted: {len(catalog)}")

    tunal_records = []
    with open('data/Caras mas recientes TUNAL.csv', 'r', encoding='latin1') as f:
        for idx, row in enumerate(csv.DictReader(f, delimiter=';')):
            prod = (row.get('Productos') or '').strip()
            cat, sub = classify(prod, row.get('Segmento', ''))
            marca = extract_tunal_brand(prod)
            try:
                caras = int(float(row.get('NumCaras', '0')))
            except:
                caras = 0
            tunal_records.append({
                'id': f"T-{idx+1}",
                'idDC': str(row.get('idDC', '')).strip(),
                'origen': 'El Tunal',
                'fecha': str(row.get('Fecha_DC', '')).strip(),
                'ruta': str(row.get('Ruta', '')).strip() or 'Sin Ruta',
                'sucCliente': str(row.get('SucCliente', '')).strip() or 'Sin Sucursal',
                'segmento': str(row.get('Segmento', '')).strip(),
                'producto': prod,
                'marca': marca,
                'categoria': cat,
                'subcategoria': sub,
                'presentacion': extract_presentation(prod),
                'caras': caras
            })

    comp_records = []
    with open('data/Caras mas recientes Competencia.csv', 'r', encoding='latin1') as f:
        for idx, row in enumerate(csv.DictReader(f, delimiter=';')):
            prod = (row.get('Productos') or '').strip()
            cat, sub = classify(prod)
            marca = extract_competitor_brand(prod)
            try:
                caras = int(float(row.get('NroCaras', '0')))
            except:
                caras = 0
            comp_records.append({
                'id': f"C-{idx+1}",
                'idDC': str(row.get('idDC', '')).strip(),
                'origen': 'Competencia',
                'fecha': str(row.get('Fecha_DC', '')).strip(),
                'ruta': str(row.get('Ruta', '')).strip() or 'Sin Ruta',
                'sucCliente': str(row.get('SucCliente', '')).strip() or 'Sin Sucursal',
                'segmento': '',
                'producto': prod,
                'marca': marca,
                'categoria': cat,
                'subcategoria': sub,
                'presentacion': extract_presentation(prod),
                'caras': caras
            })

    print(f"Total Tunal records: {len(tunal_records)}")
    print(f"Total Competencia records: {len(comp_records)}")

    # Category counts debug
    t_cat = {}
    for r in tunal_records:
        t_cat[r['categoria']] = t_cat.get(r['categoria'], 0) + 1
    print("Tunal Category Breakdown:", t_cat)

    rutas = sorted(list(set(r['ruta'] for r in tunal_records + comp_records if r['ruta'] and r['ruta'] != 'Sin Ruta')))
    sucursales = sorted(list(set(r['sucCliente'] for r in tunal_records + comp_records if r['sucCliente'] and r['sucCliente'] != 'Sin Sucursal')))
    competitor_brands = sorted(list(set(r['marca'] for r in comp_records if r['marca'])))

    data_payload = {
        'metadata': {
            'totalRecords': len(tunal_records) + len(comp_records),
            'totalTunal': len(tunal_records),
            'totalCompetencia': len(comp_records),
            'totalRutas': len(rutas),
            'totalSucursales': len(sucursales),
            'totalCompetitorBrands': len(competitor_brands),
            'rutas': rutas,
            'sucursales': sucursales,
            'competitorBrands': competitor_brands,
            'categories': list(CATEGORY_STRUCTURE.keys()),
            'categoryStructure': CATEGORY_STRUCTURE
        },
        'catalog': catalog,
        'tunalRecords': tunal_records,
        'compRecords': comp_records
    }

    with open('src/data/pricingData.json', 'w', encoding='utf-8') as f:
        json.dump(data_payload, f, ensure_ascii=False, indent=2)
    print("Saved to src/data/pricingData.json successfully!")

    import os
    os.makedirs('public/data', exist_ok=True)
    with open('public/data/pricingData.json', 'w', encoding='utf-8') as f:
        json.dump(data_payload, f, ensure_ascii=False)
    print("Saved to public/data/pricingData.json successfully!")

if __name__ == '__main__':
    main()
